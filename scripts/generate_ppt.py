from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # Helper function to add title and text
    def add_slide(prs, title_text, content_text, layout_idx=1):
        slide_layout = prs.slide_layouts[layout_idx]
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = title_text
        
        if len(slide.placeholders) > 1:
            body_shape = slide.placeholders[1]
            tf = body_shape.text_frame
            tf.text = content_text
            
            # Adjust font size for better fit
            for paragraph in tf.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(18)
            
        return slide

    # Slide 1: Title
    slide_1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide_1.shapes.title.text = "SMART INDIA HACKATHON 2026"
    slide_1.placeholders[1].text = (
        "Problem Statement ID: SIH26102\n"
        "Title: AI-Powered System to Detect Anomalies, Fraud, and Inefficiencies in MPLAD Scheme Implementation\n"
        "Theme: Smart Automation\n"
        "Category: Software\n"
        "Team: CivicShield AI"
    )

    # Slide 2: Problem & Solution
    add_slide(prs, "CivicShield AI: Explainable Risk Intelligence", 
              "REAL-WORLD ISSUE:\n"
              "MPLADS requires monitoring thousands of civil projects. Tracking financial progress, "
              "execution timelines, and inspections at scale makes it difficult to manually prioritize high-risk anomalies.\n\n"
              "SOLUTION:\n"
              "AI-assisted decision support that detects unusual project patterns, explains the contributing signals, "
              "and routes cases to authorized officials for verified field investigation.\n\n"
              "RESPONSIBLE AI PRINCIPLE:\n"
              "ANOMALY != FRAUD. AI identifies unusual patterns. Authorized officials verify the facts.")

    # Slide 3: Technical Approach
    add_slide(prs, "TECHNICAL APPROACH", 
              "METHODOLOGY:\n"
              "1. Data Ingestion (eSAKSHI & Edge records)\n"
              "2. Hybrid AI Risk Scoring (Isolation Forest + Statutory Rules)\n"
              "3. Explainability (Map statistical anomalies to guidelines)\n"
              "4. Admin Workflow (Route flagged projects for review)\n"
              "5. Field Inspection (Location-aware GPS + Photo evidence)\n"
              "6. Resolution (Immutable case tracking & Audit Logging)\n\n"
              "TECHNOLOGIES USED:\n"
              "- Frontend: React 18, Tailwind CSS, Vite\n"
              "- Backend: FastAPI, PostgreSQL, SQLAlchemy\n"
              "- ML Engine: Python, Scikit-Learn (Isolation Forest)\n"
              "- DevOps: Docker Compose, Nginx")

    # Slide 4: Feasibility and Viability
    add_slide(prs, "FEASIBILITY AND VIABILITY", 
              "DATA FEASIBILITY:\n"
              "Designed for official project records; prototype validated using synthetic data with controlled anomaly scenarios.\n\n"
              "TECHNICAL & EXECUTION FEASIBILITY:\n"
              "Built on proven, open-source stacks deployed via Docker containers. Modular 6-service architecture separated by team roles.\n\n"
              "VIABILITY (EXPLAINABLE & TAMPER-EVIDENT):\n"
              "- Rejects black-box accusations by mapping signals to MPLADS Guidelines.\n"
              "- Secures investigation lifecycle via strict Role-Based Access Control (RBAC).\n"
              "- Helps authorities focus physical inspection resources efficiently.")

    # Slide 5: Impact and Benefits
    add_slide(prs, "IMPACT AND BENEFITS", 
              "KEY IMPACTS:\n"
              "- Targeted Discovery: Highlights execution delays and financial divergence early.\n"
              "- Clear Explainability: Shows exactly which signals triggered the risk alert.\n"
              "- Verifiable Evidence: Connects digital records with ground reality via location-aware photos.\n"
              "- Demographic Tracking: Monitors compliance with SC (15%) and ST (7.5%) allocation floors.\n\n"
              "THE VALUE CHAIN:\n"
              "PROJECT DATA -> ANOMALY DETECTION -> RISK PRIORITIZATION -> TARGETED INVESTIGATION -> FIELD EVIDENCE -> ACCOUNTABILITY\n\n"
              "TRANSFORMATION:\n"
              "From untargeted manual audits to data-driven, targeted field verifications.")

    # Slide 6: Research and References
    add_slide(prs, "RESEARCH AND REFERENCES", 
              "OFFICIAL REFERENCES:\n"
              "1. MoSPI MPLADS Guidelines 2023 (Statutory Rules)\n"
              "2. eSAKSHI & PFMS TSA Architecture (Data Context)\n"
              "3. CAG Audit Reports on MPLADS (Common Procedural Deviations)\n"
              "4. Isolation Forest Algorithm (Liu et al. - Unsupervised ML)\n"
              "5. Role-Based Access Control Standards\n\n"
              "CONCLUSION:\n"
              "CivicShield AI does not replace investigators. It helps them know where to look first, "
              "why to look there, and what evidence to verify.\n\n"
              "RISK -> REASON -> RESPONSE.")

    output_path = r"c:\Users\USER\Desktop\SIH26102\CivicShield_AI_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
